# -*- coding: utf-8 -*-
"""通用量化交易百大指標 Excel: 公式表 + 數據出處 + NDX真實數據活公式示範"""
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = 'Quant100_Indicators_Formulas_2026.xlsx'
ARIAL = 'Arial'
F_T1 = Font(name=ARIAL, bold=True, size=12)
F_HDR = Font(name=ARIAL, bold=True, size=9)
F_TXT = Font(name=ARIAL, size=9)
F_GRN = Font(name=ARIAL, size=9, color='008000')
FILL_HDR = PatternFill('solid', fgColor='D9D9D9')
TIER_FILL = {
 '一':PatternFill('solid', fgColor='DDEBF7'), '二':PatternFill('solid', fgColor='E2EFDA'),
 '三':PatternFill('solid', fgColor='FFF2CC'), '四':PatternFill('solid', fgColor='FCE4D6'),
 '五':PatternFill('solid', fgColor='EDEDED'), '六':PatternFill('solid', fgColor='E4DFEC')}

# (指標, 類別, 梯隊, 計算公式, 常用參數, 訊號用法, 數據出處, 頻率)
L = [
("200日移動平均線 SMA200","趨勢","一","SMA_n = (1/n)·Σ C_{t-i}, i=0..n-1","n=200日(=40週)","C>SMA200且SMA上斜=牛市制度; 下穿防守","任何OHLCV行情源(交易所/Yahoo/Stooq/TradingView)","日"),
("12個月時序動能 12-1 Momentum","動能","一","Mom = C_{t-21} / C_{t-252} − 1 (跳過最近1個月)","回看252日, 跳21日","Mom>0做多; <0離場/做空; 月度再平衡","OHLCV行情源(月線亦可: C_{m-1}/C_{m-12}−1)","月"),
("均線交叉 50/200 (金叉/死叉)","趨勢","一","訊號 = sign(SMA50 − SMA200) 之變號","50/200日","金叉=中長期轉多; 死叉鎖定大熊市","OHLCV行情源","日"),
("MACD","趨勢/動能","一","MACD = EMA12(C) − EMA26(C); Signal = EMA9(MACD); Hist = MACD − Signal","12/26/9","Hist翻正翻負, 頂底背離","OHLCV行情源","日"),
("RSI 相對強弱","動能","一","RSI = 100 − 100/(1+RS); RS = SMMA(漲幅,14)/SMMA(跌幅,14) (Wilder平滑)","n=14","<30超賣/>70超買; 50軸多空; 背離","OHLCV行情源","日"),
("VIX 恐慌指數","波動率","一","CBOE以SPX期權全行權價鏈按方差互換公式合成之30日隱含波動率","30日期限","<20低波動制度; >25持續=風險制度","CBOE (cboe.com) / FRED: VIXCLS","日"),
("ATR 真實波幅","波動率","一","TR = max(H−L, |H−C_prev|, |L−C_prev|); ATR = SMMA(TR, 14)","n=14","止損距離=k·ATR; 倉位=風險預算/ATR","OHLCV行情源(需高低價)","日"),
("布林通道 Bollinger Bands","波動率/均值回歸","一","Mid = SMA20; Upper/Lower = Mid ± 2·σ20 (σ=收盤標準差)","20日, 2σ","帶寬擠壓後突破; 通道回歸交易","OHLCV行情源","日"),
("距52週高/低位距離","趨勢","一","DD52 = C/max(C,252日) − 1; UP52 = C/min(C,252日) − 1","252日","DD52>-6%健康; ≤-15%熊市警戒","OHLCV行情源","日/週"),
("成交量均量比","成交量","一","VR = V_t / SMA(V,20)","20日","突破日VR>1.5=有效突破確認","OHLCV行情源","日"),
("VWAP 成交量加權均價","成交量/執行","一","VWAP = Σ(TP_i·V_i)/Σ V_i, TP=(H+L+C)/3, 當日累計","日內累計","價>VWAP日內偏多; 機構執行基準","交易所tick/分鐘數據","日內"),
("EMA 指數移動平均","趨勢","一","EMA_t = α·C_t + (1−α)·EMA_{t-1}, α = 2/(n+1)","12/26/50","快慢EMA排列與交叉","OHLCV行情源","日"),
("OBV 能量潮","成交量","一","OBV_t = OBV_{t-1} + V_t·sign(C_t − C_{t-1})","累計","OBV與價背離=吸籌/派發","OHLCV行情源","日"),
("ADX 趨勢強度","趨勢強度","一","DX = 100·|DI+ − DI−|/(DI+ + DI−); ADX = SMMA(DX,14); DI±由方向移動DM±/TR算出","n=14","ADX>25用順勢策略; <20用均值回歸","OHLCV行情源(需高低價)","日"),
("KD 隨機指標","動能","一","%K = 100·(C − L14)/(H14 − L14); %D = SMA(%K, 3)","14/3/3","20/80區交叉; 高檔鈍化辨識強勢","OHLCV行情源(需高低價)","日"),
("騰落線 A/D Line","廣度","二","ADL_t = ADL_{t-1} + (漲家數 − 跌家數)","全巿場","指數新高而ADL不配合=頂部背離","NYSE/NASDAQ每日漲跌家數(交易所/Barchart/StockCharts $NYAD)","日"),
("%成份股>200日線","廣度","二","Pct = (成份股中C>SMA200之家數)/總家數","200日",">60%健康; <40%崩壞; 極低值=底部推力前奏","Barchart $NDXA200R/$SPXA200R, StockCharts","日/週"),
("淨新高-新低","廣度","二","NHNL = 52週新高家數 − 新低家數; 常用10日累計","10日累計","持續為正=EASY; 淨新低擴張=熊市確認","交易所新高新低統計(Barchart/StockCharts $NYHL)","日"),
("McClellan Oscillator/Summation","廣度","二","Osc = EMA19(淨漲家) − EMA39(淨漲家); Summation = Σ Osc","19/39","Osc±極值=超買賣; Summation<0=弱市","NYSE漲跌家數自算 / StockCharts $NYMO $NYSI","日"),
("高收益債利差 HY OAS","信用","二","指數層面期權調整利差: HY債殖利率曲線相對國債之OAS(供應商模型計算)","ICE BofA HY指數","<350bp收窄=EASY; >450bp走闊=制度轉變","FRED: BAMLH0A0HYM2 (ICE BofA)","日"),
("收益率曲線 10Y-3M / 10Y-2Y","宏觀/利率","二","Spread = Y10 − Y3M (或 Y10 − Y2)","—","倒掛(<0)=衰退領先6-18月; 倒掛後熊陡=衰退開始","FRED: T10Y3M, T10Y2Y, DGS10/DGS2/DGS3MO","日"),
("Fed政策方向與會議定價","宏觀/流動性","二","隱含利率 = 100 − FF期貨價; 會議變動機率由相鄰期貨合約反推","30天FF期貨","加息定價>50%=HARD傾向; 減息週期=EASY","CME FedWatch (cmegroup.com) / FOMC聲明","日"),
("Put/Call 比率","情緒","二","PCR = 認沽成交量 / 認購成交量 (股票型, 10日SMA)","10日均",">0.85恐慌=反向偏多; <0.45過熱","CBOE每日統計 (cboe.com/us/options/market_statistics)","日"),
("VIX期限結構","波動率","二","Ratio = VIX / VIX3M","30日/93日","<0.95 contango正常; >1.0倒掛持續=恐慌制度","CBOE VIX與VIX3M / vixcentral.com","日"),
("美元指數DXY動能","跨巿場","二","ROC26w = DXY_t / DXY_{t-130} − 1","26週","急升>+10%=全球流動性收縮警報","ICE DXY / FRED: DTWEXBGS(貿易加權)","日"),
("已實現波動率 HV","波動率","二","HV = std(ln(C_t/C_{t-1}), n) × √252","20/60日","波動率目標倉位: w = 目標σ/HV","OHLCV行情源自算","日"),
("IV Rank / IV Percentile","波動率/期權","二","IVR = (IV − min252(IV))/(max252(IV) − min252(IV))","252日窗","IVR>50賣方策略佔優; <20買方佔優","期權鏈隱含波動率(券商/ORATS/CBOE)","日"),
("銅金比","跨巿場","二","Ratio = 銅價/金價; 訊號 = Ratio之26週ROC","26週","回升=增長預期↑; 急跌=避險壓倒增長","COMEX銅+金現貨(交易所/LBMA/FRED)","日"),
("AAII/II 情緒調查","情緒","二","Spread = 牛派% − 熊派% (4週均)","4週均","<-30%極恐=反向買點; >+30%過熱","AAII (aaii.com/sentimentsurvey) / Investors Intelligence","週"),
("COT 期貨持倉報告","情緒/籌碼","二","投機者淨持倉 = 非商業多單 − 空單; 標準化: 3年Z-score","3年Z","極端擁擠反向; 商業對沖盤逆勢佈局","CFTC COT週報 (cftc.gov)","週"),
("ROC 變動率","動能","三","ROC = C_t / C_{t-n} − 1","10/20/60日","多週期ROC合成動能分數","OHLCV行情源","日"),
("Ichimoku 一目均衡表","趨勢","三","轉換=(H9+L9)/2; 基準=(H26+L26)/2; 先行A=(轉+基)/2前移26; 先行B=(H52+L52)/2前移26","9/26/52","價在雲上+轉換>基準=多頭配置","OHLCV行情源(需高低價)","日"),
("Donchian 通道","趨勢/突破","三","Upper = max(H, n); Lower = min(L, n)","20日(海龜)","突破n日高做多, 破n/2日低離場","OHLCV行情源","日"),
("Keltner 通道","波動率","三","Mid = EMA20; 帶 = Mid ± 2·ATR10","20/10/2","布林縮進Keltner內=擠壓(Squeeze)","OHLCV行情源","日"),
("拋物線 SAR","趨勢","三","SAR_t = SAR_{t-1} + AF·(EP − SAR_{t-1}); AF由0.02遞增至0.2","0.02/0.2","移動止損; 價穿SAR=反轉","OHLCV行情源","日"),
("SuperTrend","趨勢","三","ST = (H+L)/2 ± m·ATR(p) (帶方向鎖定規則)","p=10, m=3","價在ST上=持多; 下穿=離場","OHLCV行情源","日"),
("Williams %R","動能","三","%R = −100·(H14 − C)/(H14 − L14)","n=14","<-80超賣; >-20超買","OHLCV行情源","日"),
("CCI 順勢指標","動能","三","CCI = (TP − SMA20(TP))/(0.015·平均絕對偏差), TP=(H+L+C)/3","n=20","±100突破順勢進場","OHLCV行情源","日"),
("MFI 資金流量指數","量價","三","MF = TP·V; MFI = 100 − 100/(1 + 正MF14/負MF14)","n=14","帶量RSI; 20/80超買賣+背離","OHLCV行情源","日"),
("Chaikin Money Flow","量價","三","CMF = Σ21[MFM·V]/Σ21V; MFM = ((C−L)−(H−C))/(H−L)","n=21","CMF>0資金流入; 與價背離預警","OHLCV行情源","日"),
("Accumulation/Distribution","量價","三","ADL_t = ADL_{t-1} + MFM·V (MFM同CMF)","累計","與OBV互證吸籌派發","OHLCV行情源","日"),
("Force Index 力度指數","量價","三","FI = EMA13((C_t − C_{t-1})·V_t)","n=13","翻正翻負+背離","OHLCV行情源","日"),
("TRIX","趨勢","三","TRIX = 1期ROC( EMA(EMA(EMA(C,15),15),15) )","n=15","三重平滑濾噪; 零軸與訊號線交叉","OHLCV行情源","日"),
("Aroon","趨勢","三","AroonUp = 100·(n − 距n日最高日數)/n; Down同理用最低","n=25","Up>70且Down<30=強趨勢; 交叉=轉勢","OHLCV行情源","日"),
("Ultimate Oscillator","動能","三","UO = 100·(4·Avg7 + 2·Avg14 + 1·Avg28)/7; Avg=ΣBP/ΣTR, BP=C−min(L,C_prev)","7/14/28","三週期防假背離","OHLCV行情源","日"),
("Vortex 指標","趨勢","三","VI+ = Σ14|H_t − L_{t-1}|/Σ14TR; VI− = Σ14|L_t − H_{t-1}|/Σ14TR","n=14","VI+上穿VI−=轉多","OHLCV行情源","日"),
("Pivot Points 樞軸點","支撐阻力","三","P = (H+L+C)/3; R1 = 2P − L; S1 = 2P − H; R2 = P + (H−L)...","前日/前週數據","日內支撐阻力網格","OHLCV行情源","日內"),
("斐波那契回撤","支撐阻力","三","回撤位 = 高點 − (高點−低點)×{0.382, 0.5, 0.618}","38.2/50/61.8%","趨勢回調進場參考位","OHLCV行情源(擺動高低點)","任意"),
("Heikin-Ashi 平均K線","趨勢","三","HA_C = (O+H+L+C)/4; HA_O = (HA_O_prev + HA_C_prev)/2","—","連續同色HA蠟燭=趨勢延續","OHLCV行情源","日"),
("波段結構 HH-HL/LH-LL","價格行為","三","Swing高/低 = 前後k日之極值; 序列HH+HL=升勢, LH+LL=跌勢","k=5-10","道氏理論量化版; 結構破壞=轉勢","OHLCV行情源","日"),
("M2 貨幣供應年增率","流動性","四","YoY = M2_t / M2_{t-12} − 1","12月","回升且>4%=寬鬆背景; 負增長=罕見緊縮","FRED: M2SL (月, 聖路易聯儲)","月"),
("淨流動性","流動性","四","NetLiq = Fed總資產(WALCL) − 財政部TGA(WTREGEN) − 逆回購(RRPONTSYD)","週變化","按週回升=EASY; 收縮加速=HARD","FRED: WALCL, WTREGEN, RRPONTSYD","週"),
("NFCI 金融條件指數","流動性","四","芝加哥聯儲105項金融指標動態因子模型加權(0=歷史均值)","—","<-0.3寬鬆; >+0.3且上行=收緊","FRED: NFCI (芝加哥聯儲)","週"),
("核心CPI年增率","通脹","四","YoY = CPI核心_t / CPI核心_{t-12} − 1","12月","<4%或高位回落=政策空間開; >6%上行=緊縮壓力","FRED: CPILFESL (BLS)","月"),
("ISM 製造業PMI","宏觀","四","5分項擴散指數等權: 新訂單/生產/就業/交付/庫存, 各=P(改善)+0.5·P(持平)","50=榮枯線",">52回升=擴張; <45=收縮深化","ISM官網 (ismworld.org); 歷史見FRED/宏觀供應商","月"),
("Sahm 規則失業率動能","宏觀","四","Sahm = U3三月均 − 過去12月U3三月均之最低值","0.5pt閾值","≥+0.50pt=衰退已開始(實時判定金標準)","FRED: SAHMREALTIME (BLS U3自算亦可)","月"),
("初領失業救濟金","宏觀","四","4週移動平均; 訊號 = 相對26週低點之變幅","4週均","+15%以上且上行=勞動巿場轉弱","FRED: ICSA (勞工部, 週四發布)","週"),
("LEI 領先經濟指標","宏觀","四","10項領先成份加權合成; 訊號 = 6個月變率年化","6月年化","<-4%且擴散差=衰退警戒","Conference Board (conference-board.org)","月"),
("BAA-AAA 投資級利差","信用","四","Spread = Moody's BAA殖利率 − AAA殖利率","—","<0.95%=EASY; >1.35%走闊=信用壓力(1919年起百年序列)","FRED: BAA, AAA (Moody's)","月/日"),
("商業票據/短期融資利差","信用/流動性","四","CP利差 = 3月非金融CP利率 − 3月T-bill (TED已停: 改用SOFR−T-bill)","3個月期",">50bp=企業短融凍結(2008/2020型)","FRED: CPN3M, TB3MS, SOFR","日"),
("SLOOS 銀行信貸標準","信用","四","淨收緊% = 收緊銀行比例 − 放鬆比例 (工商貸款)","季調查",">+20%大幅收緊=領先信貸緊縮1-2季","Fed SLOOS (federalreserve.gov, 季度)","季"),
("油價年變率","商品/通脹","四","ROC12m = Brent_t / Brent_{t-12m} − 1","12月","> +80%=油震盪衰退前兆(1973/79/90/22)","FRED: DCOILBRENTEU, DCOILWTICO / EIA","日"),
("黃金相對強度","商品/避險","四","RS = (Gold/SPX)之13週ROC","13週","持續走強=避險需求/貨幣信心弱","LBMA金價+指數行情自算","日"),
("EPS 修正廣度","盈利","四","Breadth = 上調家數/(上調+下調) (3月均)","3月均",">55%=盈利動能佳; <45%下行=盈利衰退前奏","LSEG(Refinitiv)/FactSet Earnings Revisions","週"),
("前瞻P/E","估值","四","FwdPE = P / 未來12月共識EPS","NTM","相對自身5年分位判貴賤; >90分位+利率上行=危險","FactSet Earnings Insight(免費週報)/LSEG","週"),
("股權風險溢價 ERP","估值","四","ERP = 1/FwdPE − 10Y名義殖利率","—",">3%股票吸引; <0股票貴過債券","自算: FactSet PE + FRED DGS10","月"),
("Shiller CAPE","估值","四","CAPE = 實質價格 / 過去10年實質EPS平均","10年","相對自身10年中位; 泡沫帶=未來10年報酬受壓","Robert Shiller數據集 (econ.yale.edu/~shiller/data.htm)","月"),
("股息率與股息增長","估值/現金流","四","DY = 過去12月股息/價格; DG = 股息YoY","TTM","股息負增長=現金流壓力訊號","S&P DJI指數股息數據/供應商","季"),
("巴菲特指標","估值","四","Ratio = 美股總巿值 / 名義GDP; 訊號=相對自身5年均值偏離","5年均","歷史極端高位=長週期報酬受壓","FRED: 自算(Wilshire 5000×係數 / GDP)","季"),
("滾動 Beta","風險","四","β = Cov(r_i, r_m)/Var(r_m), 滾動252日","252日","組合巿場暴露監控與對沖比率","行情報酬序列自算","日"),
("買賣價差 Bid-Ask Spread","微結構","五","Spread% = (Ask − Bid)/中間價","L1報價","價差急闊=流動性抽走預警","交易所L1報價(券商API/TAQ數據)","tick"),
("訂單簿失衡 OBI","微結構","五","OBI = (ΣBidVol − ΣAskVol)/(ΣBidVol + ΣAskVol), 前5檔","L2前5檔",">0.3短線偏多; 高頻方向預測因子","交易所L2訂單簿(券商API)","tick"),
("訂單流 Delta / CVD","微結構","五","Delta = 主動買量 − 主動賣量(按成交發生於買/賣價判定); CVD = 累計Delta","逐筆","CVD與價背離=吸收; 突破配合Delta確認","交易所逐筆成交(tick data)","tick"),
("大單/區塊交易追蹤","微結構/籌碼","五","篩選單筆成交額 > 閾值(如$1M)之淨方向","自定閾值","機構足跡; 暗池印花追蹤","TAQ/交易所tick, FINRA暗池週報(ATS)","日"),
("伽瑪暴露 GEX","期權結構","五","GEX = Σ(Γ_i × OI_i × 100 × S²) × 做巿商方向假設(Call+/Put−)","全期權鏈","正伽瑪=波動被壓; 深負伽瑪=追漲殺跌放大","期權鏈OI+Greeks(CBOE/ORATS/SqueezeMetrics)","日"),
("期權未平倉 OI 與 Max Pain","期權結構","五","MaxPain = argmin_K Σ(期權買方到期內在值)","到期鏈","到期日價格磁吸參考; OI牆=支撐阻力","期權鏈數據(CBOE/券商)","日"),
("期現基差 Basis","衍生品","五","Basis% = (F − S)/S × (365/到期天數)","近月合約","基差極端=情緒/資金成本極端","期貨+現貨行情(CME等)","日"),
("資金費率 Funding Rate","衍生品/加密","五","每8小時費率 = clamp(溢價指數, ±上限); 年化≈費率×3×365","永續合約","極高正費率=多頭擁擠反向訊號","交易所API(Binance/OKX/Bybit)","8小時"),
("期限結構 Contango/Backwardation","衍生品/商品","五","Slope = (F2 − F1)/F1 (次月−近月)","近兩月合約","Backwardation=現貨緊張; 展期收益方向","期貨鏈行情(CME/ICE)","日"),
("融券餘額 Short Interest","籌碼","五","SI% = 融券股數/流通股本; DTC = 融券股數/日均量","雙週數據","SI高+突破=軋空燃料; DTC>5天危險","交易所雙週SI報告(NASDAQ/NYSE)/FINRA","雙週"),
("融資餘額 Margin Debt","槓桿","五","YoY = MarginDebt_t / MarginDebt_{t-12} − 1","12月","> +50%過熱; 頂部前融資往往先見頂","FINRA月度統計 (finra.org margin statistics)","月"),
("基金資金流","資金流","五","4週淨流入 / 資產規模","4週累計","極端流出=反向底部訊號","ICI週報 (ici.org) / EPFR(付費)","週"),
("內部人買賣比","籌碼","五","Ratio = 內部人買入金額 / 賣出金額 (全巿場, 月度)","月度","急升>1=大跌後聰明錢進場(反向)","SEC Form 4 (EDGAR) / OpenInsider匯總","週/月"),
("季節性效應","統計","五","同月平均報酬與勝率: mean(r_月m), P(r_月m>0), 排除重疊樣本","30年+樣本","11-4月強於5-10月; 選舉週期第3年最強","歷史行情自算(需長樣本)","月"),
("新聞/社交情緒 NLP","另類數據","五","情緒分數 = (正面詞/文章 − 負面)/總數 或 LLM評分之滾動Z-score","7日Z","極端負面+價格鈍化=利空出盡","GDELT(免費)/RavenPack(付費)/X-API","日"),
("Z-score 偏離度","統計套利","六","Z = (P − SMA_n)/σ_n (單資產) 或 spread之Z (配對)","n=20-60","|Z|>2進場回歸, Z→0平倉","行情自算","日"),
("協整檢驗","統計套利","六","Engle-Granger: 迴歸P_A=α+β·P_B+ε, 對ε做ADF單位根檢驗 p<0.05","2-3年窗","協整對→配對交易; 破裂監控","行情自算(statsmodels)","日"),
("Hurst 指數","統計","六","R/S分析: E[R(n)/S(n)] ∝ n^H, 取log-log斜率","H閾0.5","H>0.55趨勢市用順勢; <0.45用均值回歸","行情自算","任意"),
("GARCH 波動率預測","統計","六","GARCH(1,1): σ²_t = ω + α·ε²_{t-1} + β·σ²_{t-1}","(1,1)","事前波動率預測→倉位/期權定價","報酬序列自算(arch套件)","日"),
("Kalman 濾波","統計","六","狀態空間: β_t = β_{t-1}+w; y_t = β_t·x_t+v; 遞歸更新後驗","Q/R自校","時變對沖比率/動態α估計","行情自算(pykalman等)","日"),
("因子暴露分析","因子","六","時序迴歸: r_i − rf = α + β_mkt·MKT + β_smb·SMB + β_hml·HML + β_mom·MOM + ε","FF3/FF5+動能","組合風格歸因; 非意圖暴露對沖","Ken French Data Library(免費, dartmouth.edu)","月"),
("滾動相關性矩陣","風控","六","ρ_ij = Cov(r_i,r_j)/(σ_i·σ_j), 滾動60日","60日","分散度監控; 危機時相關性→1","報酬序列自算","日"),
("夏普比率 Sharpe","績效","六","Sharpe = (R_p − R_f)/σ_p × √252 (日頻年化)","年化",">1可用; >2優秀; 比較策略標準","策略淨值序列自算; R_f用FRED: DGS3MO","任意"),
("索提諾比率 Sortino","績效","六","Sortino = (R_p − R_f)/下行偏差σ_d × √252","年化","只罰下行波動; 適合偏態策略","策略淨值自算","任意"),
("最大回撤 MaxDD","風控","六","MaxDD = max_t(1 − Equity_t / max_{s≤t}Equity_s)","全樣本","生存能力量度; 決定槓桿上限","策略淨值自算","任意"),
("卡瑪比率 Calmar","績效","六","Calmar = CAGR / |MaxDD| (常用近3年)","3年窗",">1合格; 趨勢CTA常用","策略淨值自算","任意"),
("資訊比率 IR","績效","六","IR = (R_p − R_bench)/TrackingError, TE = std(超額報酬)×√252","年化","主動管理超額報酬穩定性","策略+基準淨值自算","任意"),
("勝率與盈虧比","績效","六","WinRate = 獲利筆數/總筆數; Payoff = 平均獲利/|平均虧損|","逐筆","期望值 = WR×Payoff − (1−WR); 兩者需合看","交易紀錄自算","逐筆"),
("盈虧因子 Profit Factor","績效","六","PF = Σ獲利 / |Σ虧損|","逐筆",">1.5合格; >3優秀","交易紀錄自算","逐筆"),
("凱利公式 Kelly %","倉位管理","六","f* = W − (1−W)/R, W=勝率, R=盈虧比; 實務用半凱利f*/2","半凱利","由統計優勢推最優倉位比例","交易紀錄自算","逐筆"),
]
assert len(L) == 100, len(L)

wb = openpyxl.Workbook()

# ============ 00 說明與數據出處 ============
ws = wb.active; ws.title = '00_說明與數據出處'
rows0 = [
("通用量化交易百大指標 — 計算公式與數據出處 (2026-07)", ""),
("", ""),
("排列", "按重要性由高至低, 序號越低越重要。梯隊: 一=核心價格/趨勢/波動(1-15) 二=廣度/情緒/跨巿場(16-30) 三=進階技術(31-50) 四=宏觀/流動性/基本面(51-70) 五=微結構/衍生品/另類(71-85) 六=統計/因子/績效風控(86-100)。"),
("符號約定", "C=收盤 O=開盤 H=最高 L=最低 V=成交量 P=價格 F=期貨價 S=現貨價 r=報酬 σ=標準差 n=回看期。SMA=簡單移動平均; EMA=指數移動平均(α=2/(n+1)); SMMA=Wilder平滑(α=1/n); ROC=變動率; TP=(H+L+C)/3典型價; TR=真實波幅; OI=未平倉合約。"),
("主要免費數據出處", "① FRED (fred.stlouisfed.org): 利率/利差/宏觀/流動性全系列, 表內列明代碼 | ② CBOE (cboe.com): VIX/VIX3M/SKEW/Put-Call | ③ CFTC (cftc.gov): COT持倉週報 | ④ FINRA (finra.org): 融資餘額/暗池 | ⑤ SEC EDGAR: Form 4內部人 | ⑥ ICI (ici.org): 基金資金流 | ⑦ Conference Board: LEI/消費者信心 | ⑧ ISM (ismworld.org): PMI | ⑨ Shiller (econ.yale.edu/~shiller): CAPE百年數據 | ⑩ Ken French Data Library: 因子報酬 | ⑪ OHLCV行情: 交易所/Yahoo Finance/Stooq/TradingView | ⑫ 付費: LSEG/FactSet(盈利修正), EPFR(資金流), RavenPack(NLP), ORATS(期權)。"),
("使用要點", "① 單一指標勝率有限, 價值在跨類別組合確認(如趨勢+廣度+信用三重確認); ② 1-30名決定大方向, 31名後偏輔助/特定策略; ③ 86-100名不產生買賣訊號, 但決定策略能否存活; ④ 參數為業界慣用預設, 實盤前應對自身標的重新校準並做樣本外檢驗。"),
("工作表", "01_百大指標公式表 = 100項完整定義 | 02_計算示範_NDX = 用原工作簿2016-2026真實NDX日線, 以活Excel公式示範12個代表指標的逐格計算(可展開查看每條公式)。"),
("與前次交付的關係", "前次 AFable...R5.0.xlsx 之『10_百大指標庫』是專為Easy/Hard Money制度判區設計的宏觀制度指標庫; 本檔是通用量化交易常用指標總表, 兩者互補。"),
]
for i, (a, b) in enumerate(rows0, 1):
    ws.cell(row=i, column=1, value=a).font = F_T1 if i == 1 else Font(name=ARIAL, bold=True, size=9)
    c = ws.cell(row=i, column=2, value=b); c.font = F_TXT
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 150

# ============ 01 百大指標公式表 ============
ws = wb.create_sheet('01_百大指標公式表')
ws.cell(row=1, column=1, value='通用量化交易百大指標 | 重要性由高至低(序號越低越重要) | 公式符號見00表 | 數據出處含FRED代碼/官網').font = F_T1
hdr = ['序號','指標名稱','類別','梯隊','計算公式','常用參數','訊號用法','數據出處','頻率']
for c, h in enumerate(hdr, 1):
    cell = ws.cell(row=2, column=c, value=h); cell.font = F_HDR; cell.fill = FILL_HDR
for i, r in enumerate(L, 1):
    row = 2 + i
    vals = [i] + [r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v); cell.font = F_TXT
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=row, column=4).fill = TIER_FILL[r[2]]
for c, w in enumerate([5, 26, 12, 6, 52, 13, 34, 40, 7], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'A3'

# ============ 02 計算示範 (NDX 真實數據, 活公式) ============
nd = pd.read_csv('data/workbook_daily.csv')[['date', 'ndx']].dropna()
ws = wb.create_sheet('02_計算示範_NDX')
ws.cell(row=1, column=1, value='12個代表指標活公式示範 | 數據=原工作簿02A之NDX真實收盤(2016-07-08..2026-07-10, 2515日) | 每格公式可點開查看 | 對應01表序號標於第2行 | 收盤價only, 需H/L/V之指標(ATR/KD等)見01表公式').font = F_T1
hdr2 = ['日期','NDX收盤','日報酬','SMA50','SMA200','EMA12','EMA26','MACD','Signal9','Hist',
        '漲幅U','跌幅D','AvgU14','AvgD14','RSI14','布林上軌','布林下軌','ROC20','HV20年化%','12-1動能','距52週高','Z-score20']
maps  = ['—','數據','—','#1/#3','#1','#12','#12','#4','#4','#4','#5','#5','#5','#5','#5','#8','#8','#31','#26','#2','#9','#86']
for c, h in enumerate(hdr2, 1):
    cell = ws.cell(row=2, column=c, value=h); cell.font = F_HDR; cell.fill = FILL_HDR
    ws.cell(row=3, column=c, value=maps[c-1]).font = F_GRN
R0 = 4
n = len(nd)
for i, (_, rw) in enumerate(nd.iterrows()):
    r = R0 + i
    ws.cell(row=r, column=1, value=str(rw['date'])[:10]).font = F_TXT
    ws.cell(row=r, column=2, value=round(float(rw['ndx']), 2)).font = F_TXT
    if i >= 1:
        ws.cell(row=r, column=3, value=f'=B{r}/B{r-1}-1').font = F_TXT
        ws.cell(row=r, column=11, value=f'=MAX(B{r}-B{r-1},0)').font = F_TXT
        ws.cell(row=r, column=12, value=f'=MAX(B{r-1}-B{r},0)').font = F_TXT
    if i >= 49:
        ws.cell(row=r, column=4, value=f'=AVERAGE(B{r-49}:B{r})').font = F_TXT
    if i >= 199:
        ws.cell(row=r, column=5, value=f'=AVERAGE(B{r-199}:B{r})').font = F_TXT
    # EMA12 seed at i=11 (SMA of first 12), recursive after
    if i == 11:
        ws.cell(row=r, column=6, value=f'=AVERAGE(B{R0}:B{r})').font = F_TXT
    elif i > 11:
        ws.cell(row=r, column=6, value=f'=B{r}*2/13+F{r-1}*11/13').font = F_TXT
    if i == 25:
        ws.cell(row=r, column=7, value=f'=AVERAGE(B{R0}:B{r})').font = F_TXT
    elif i > 25:
        ws.cell(row=r, column=7, value=f'=B{r}*2/27+G{r-1}*25/27').font = F_TXT
    if i >= 25:
        ws.cell(row=r, column=8, value=f'=F{r}-G{r}').font = F_TXT
    if i == 33:
        ws.cell(row=r, column=9, value=f'=AVERAGE(H{r-8}:H{r})').font = F_TXT
    elif i > 33:
        ws.cell(row=r, column=9, value=f'=H{r}*2/10+I{r-1}*8/10').font = F_TXT
    if i >= 33:
        ws.cell(row=r, column=10, value=f'=H{r}-I{r}').font = F_TXT
    # Wilder RSI: seed avg at i=14
    if i == 14:
        ws.cell(row=r, column=13, value=f'=AVERAGE(K{R0+1}:K{r})').font = F_TXT
        ws.cell(row=r, column=14, value=f'=AVERAGE(L{R0+1}:L{r})').font = F_TXT
    elif i > 14:
        ws.cell(row=r, column=13, value=f'=(M{r-1}*13+K{r})/14').font = F_TXT
        ws.cell(row=r, column=14, value=f'=(N{r-1}*13+L{r})/14').font = F_TXT
    if i >= 14:
        ws.cell(row=r, column=15, value=f'=IF(N{r}=0,100,100-100/(1+M{r}/N{r}))').font = F_TXT
    if i >= 19:
        ws.cell(row=r, column=16, value=f'=AVERAGE(B{r-19}:B{r})+2*STDEV(B{r-19}:B{r})').font = F_TXT
        ws.cell(row=r, column=17, value=f'=AVERAGE(B{r-19}:B{r})-2*STDEV(B{r-19}:B{r})').font = F_TXT
        ws.cell(row=r, column=19, value=f'=STDEV(C{r-19}:C{r})*SQRT(252)*100').font = F_TXT
        ws.cell(row=r, column=22, value=f'=(B{r}-AVERAGE(B{r-19}:B{r}))/STDEV(B{r-19}:B{r})').font = F_TXT
    if i >= 20:
        ws.cell(row=r, column=18, value=f'=B{r}/B{r-20}-1').font = F_TXT
    if i >= 251:
        ws.cell(row=r, column=20, value=f'=B{r-21}/B{r-251}-1').font = F_TXT
        ws.cell(row=r, column=21, value=f'=B{r}/MAX(B{r-251}:B{r})-1').font = F_TXT
for c, w in enumerate([11,10,9,10,10,10,10,9,9,8,8,8,9,9,8,11,11,8,10,9,9,9], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
for col, fmt in [(3,'0.00%'),(18,'0.00%'),(20,'0.00%'),(21,'0.00%')]:
    for r in range(R0, R0+n):
        ws.cell(row=r, column=col).number_format = fmt
ws.freeze_panes = 'C4'

wb.save(OUT)
print('saved', OUT, '| rows:', n)
